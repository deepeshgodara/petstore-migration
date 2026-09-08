#!/usr/bin/env python3
"""
Approximate visual renderer for a .pptx, so layout can be eyeballed without
PowerPoint/LibreOffice. Draws shape rects + wrapped text with real sizes.
Not pixel-exact; good enough to catch collisions, overflow and bad geometry.

Usage: python3 scripts/preview_deck.py <deck.pptx> <outdir> [scale_px_per_inch]
"""
import os, sys
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

EMU = 914400
FONTS = [
    "/System/Library/Fonts/Helvetica.ttc",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
]
MONOS = [
    "/System/Library/Fonts/Menlo.ttc",
    "/System/Library/Fonts/Supplemental/Courier New.ttf",
]

def _font(path_list, px, bold=False):
    for p in path_list:
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, px, index=(1 if bold and p.endswith(".ttc") else 0))
            except Exception:
                try: return ImageFont.truetype(p, px)
                except Exception: continue
    return ImageFont.load_default()

def rgb(c, default=(200,200,200)):
    try:
        if c is None or c.rgb is None: return default
        v = str(c.rgb)
        return tuple(int(v[i:i+2],16) for i in (0,2,4))
    except Exception:
        return default

def wrap(draw, text, font, maxw):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur+" "+w).strip()
        if draw.textlength(t, font=font) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur); cur = w
    if cur: lines.append(cur)
    return lines or [""]

def render(pptx_path, outdir, ppi=110):
    prs = Presentation(pptx_path)
    W = int(prs.slide_width/EMU*ppi); H = int(prs.slide_height/EMU*ppi)
    os.makedirs(outdir, exist_ok=True)
    out = []
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W,H), (13,13,13))
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            L = int((sh.left or 0)/EMU*ppi); T = int((sh.top or 0)/EMU*ppi)
            Wd = int((sh.width or 0)/EMU*ppi); Ht = int((sh.height or 0)/EMU*ppi)
            if Wd <= 0 or Ht <= 0: continue
            # picture
            if sh.shape_type == 13:
                try:
                    im2 = Image.open(io_bytes(sh)).convert("RGB").resize((Wd,Ht))
                    img.paste(im2,(L,T))
                except Exception:
                    d.rectangle([L,T,L+Wd,T+Ht], fill=(40,40,38), outline=(90,90,88))
                    d.text((L+8,T+8), "[image]", fill=(160,160,158), font=_font(FONTS,13))
                continue
            # fill
            try:
                if sh.fill.type is not None and sh.fill.type == 1:
                    d.rectangle([L,T,L+Wd,T+Ht], fill=rgb(sh.fill.fore_color,(26,26,25)))
            except Exception: pass
            # outline
            try:
                if sh.line.fill.type == 1:
                    d.rectangle([L,T,L+Wd,T+Ht], outline=rgb(sh.line.color,(44,44,42)), width=1)
            except Exception: pass
            # table
            if getattr(sh,"has_table",False) and sh.has_table:
                tbl=sh.table; rows=len(tbl.rows); cols=len(tbl.columns)
                ch=Ht/max(1,rows)
                for r in range(rows):
                    cw=Wd/max(1,cols)
                    for c in range(cols):
                        x0=L+int(c*cw); y0=T+int(r*ch)
                        cell=tbl.cell(r,c)
                        d.rectangle([x0,y0,x0+int(cw),y0+int(ch)],
                                    fill=(36,36,34) if r==0 else (26,26,25), outline=(44,44,42))
                        p=cell.text_frame.paragraphs[0]
                        px=max(8,int((p.font.size.pt if p.font.size else 10)*ppi/72))
                        f=_font(FONTS,px,bold=bool(p.font.bold))
                        col=rgb(p.font.color,(195,194,183)) if p.font.color and p.font.color.type is not None else (195,194,183)
                        for i,ln in enumerate(wrap(d,cell.text,f,cw-14)[:2]):
                            d.text((x0+7,y0+5+i*(px+2)), ln, fill=col, font=f)
                continue
            # text
            if sh.has_text_frame and sh.text_frame.text.strip():
                y=T+4
                for p in sh.text_frame.paragraphs:
                    if not (p.text or "").strip():
                        y+=6; continue
                    r0=p.runs[0] if p.runs else None
                    pt=(r0.font.size.pt if (r0 and r0.font.size) else (p.font.size.pt if p.font.size else 12))
                    bold=bool((r0.font.bold if r0 else None) or p.font.bold)
                    name=(r0.font.name if r0 else None) or ""
                    px=max(7,int(pt*ppi/72))
                    f=_font(MONOS if name in ("Menlo","Consolas","Courier New") else FONTS, px, bold)
                    col=(255,255,255)
                    try:
                        src=(r0.font.color if r0 and r0.font.color and r0.font.color.type is not None else p.font.color)
                        col=rgb(src,(255,255,255))
                    except Exception: pass
                    for ln in wrap(d,p.text,f,max(10,Wd-10)):
                        if y > T+Ht+40: break
                        d.text((L+5,y), ln, fill=col, font=f); y += px+3
                    y += int((p.space_after.pt if p.space_after else 4)*ppi/72)
        d.text((W-70,H-22), f"{idx}", fill=(120,120,118), font=_font(FONTS,13))
        p=os.path.join(outdir,f"slide_{idx:02d}.png"); img.save(p); out.append(p)
    return out

def io_bytes(shape):
    import io
    return io.BytesIO(shape.image.blob)

if __name__=="__main__":
    files=render(sys.argv[1], sys.argv[2], int(sys.argv[3]) if len(sys.argv)>3 else 110)
    print(f"rendered {len(files)} slides -> {sys.argv[2]}")
